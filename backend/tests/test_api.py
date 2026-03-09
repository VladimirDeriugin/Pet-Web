"""Integration tests for the FastAPI document analysis endpoints."""

import io
import pytest
from httpx import AsyncClient, ASGITransport

from app.main import app


@pytest.mark.asyncio
async def test_health_endpoint():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        response = await client.get("/api/documents/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


@pytest.mark.asyncio
async def test_analyze_txt_document():
    text = (
        "Artificial intelligence is changing the world. "
        "Machine learning allows computers to learn from data. "
        "Natural language processing enables understanding of human language. "
        "Deep learning is a subset of machine learning that uses neural networks. "
        "AI applications include healthcare, finance, and transportation. "
        "The future of AI is bright and full of possibilities."
    )
    file_content = text.encode("utf-8")
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post(
            "/api/documents/analyze",
            files={"file": ("test_document.txt", io.BytesIO(file_content), "text/plain")},
        )
    assert response.status_code == 200
    data = response.json()
    assert data["filename"] == "test_document.txt"
    assert data["file_type"] == "TXT"
    assert data["stats"]["word_count"] > 0
    assert isinstance(data["keywords"], list)
    assert isinstance(data["summary"], str)
    assert data["sentiment"] in ("positive", "negative", "neutral")
    assert "language" in data
    assert isinstance(data["entities"], list)


@pytest.mark.asyncio
async def test_analyze_empty_file():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post(
            "/api/documents/analyze",
            files={"file": ("empty.txt", io.BytesIO(b""), "text/plain")},
        )
    assert response.status_code == 400


@pytest.mark.asyncio
async def test_analyze_unsupported_format():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post(
            "/api/documents/analyze",
            files={"file": ("doc.xlsx", io.BytesIO(b"fake data"), "application/vnd.ms-excel")},
        )
    assert response.status_code == 415


@pytest.mark.asyncio
async def test_analyze_docx_document():
    import docx as python_docx

    doc = python_docx.Document()
    doc.add_paragraph("This is a test DOCX document for analysis.")
    doc.add_paragraph("It contains multiple paragraphs with meaningful content.")
    doc.add_paragraph("The document analysis service should extract this text correctly.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post(
            "/api/documents/analyze",
            files={"file": ("test.docx", buf, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
    assert response.status_code == 200
    data = response.json()
    assert data["file_type"] == "DOCX"
    assert data["stats"]["word_count"] > 0
