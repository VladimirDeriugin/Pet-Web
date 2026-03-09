"""Unit tests for document_parser service."""

import io
import pytest
from app.services.document_parser import extract_text, _extract_text


class TestExtractText:
    def test_plain_text_utf8(self):
        content = "Hello world. This is a test document.".encode("utf-8")
        text, ftype = extract_text("sample.txt", content)
        assert "Hello world" in text
        assert ftype == "TXT"

    def test_plain_text_markdown(self):
        content = "# Title\n\nSome content here.".encode("utf-8")
        text, ftype = extract_text("README.md", content)
        assert "Title" in text
        assert ftype == "TXT"

    def test_unsupported_extension(self):
        with pytest.raises(ValueError, match="Unsupported file type"):
            extract_text("document.xyz", b"data")

    def test_empty_filename(self):
        # no extension → unsupported
        with pytest.raises(ValueError):
            extract_text("nodotfile", b"data")

    def test_latin1_fallback(self):
        # Bytes that are valid latin-1 but not utf-8
        content = "Caf\xe9 au lait".encode("latin-1")
        text = _extract_text(content)
        assert "Caf" in text

    def test_large_text(self):
        content = ("word " * 10_000).encode("utf-8")
        text, _ = extract_text("large.txt", content)
        assert len(text) > 1000


class TestExtractDocx:
    def test_docx_extraction(self):
        """Create a minimal DOCX in-memory and extract its text."""
        import docx
        import io as _io

        doc = docx.Document()
        doc.add_paragraph("This is the first paragraph.")
        doc.add_paragraph("This is the second paragraph.")
        buf = _io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        content = buf.read()

        text, ftype = extract_text("test.docx", content)
        assert "first paragraph" in text
        assert ftype == "DOCX"
